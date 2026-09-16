import hashlib
import io
import json
import struct
import sqlite3
from pathlib import Path
import pytest
from docx import Document
from openpyxl import Workbook
from app import db
from app.config import settings
from app.index_document import index_document
from app.worker import claim
from app.retrieval import search
from app.machine import inspect_gguf
from app.inference import runtime
from app.document_chat import summarise
from conftest import sign_in
from test_application import pdf_bytes, indexed, fake_model


def add_document(admin,kb,name,payload):
    response=admin.post(f'/api/knowledge/{kb}/documents',content=payload,
                        headers={'X-Filename':name,'Content-Type':'application/octet-stream'})
    assert response.status_code==202,response.text
    identifier=response.json()['id']
    assert claim()==identifier
    index_document(identifier)
    return identifier


def gguf():
    key=b'general.architecture';value=b'llama'
    return b'GGUF'+struct.pack('<IQQ',3,1,1)+struct.pack('<Q',len(key))+key+struct.pack('<I',8)+struct.pack('<Q',len(value))+value+b'\0'*32


def add_member(admin,kb,permission='reader'):
    admin.post('/api/users',json={'username':'member','password':'member-password-123'})
    response=admin.post(f'/api/knowledge/{kb}/members',json={'username':'member','permission':permission})
    assert response.status_code==200,response.text


@pytest.mark.parametrize('kind',['txt','md','csv','docx','xlsx'])
def test_new_formats_and_source_locations(admin,kb,kind):
    if kind in ('txt','md'):
        payload=b'Nelsonict teaches MikroTik networking.\nPractical training is available.'
    elif kind=='csv':
        payload=b'Course,Provider\nMikroTik,Nelsonict\n'
    elif kind=='docx':
        stream=io.BytesIO();doc=Document();doc.add_paragraph('Nelsonict teaches MikroTik networking.');doc.save(stream);payload=stream.getvalue()
    else:
        stream=io.BytesIO();book=Workbook();book.active.title='Courses';book.active.append(['MikroTik','Nelsonict']);book.save(stream);payload=stream.getvalue()
    identifier=add_document(admin,kb,'training.'+kind,payload)
    result=search(1,kb,'MikroTik')
    assert result and result[0]['format']==kind and result[0]['location']
    preview=admin.get(f'/api/documents/{identifier}/preview?page={result[0]["page"]}')
    assert preview.status_code==200 and 'MikroTik' in preview.json()['text']
    assert admin.get(f'/api/documents/{identifier}/download').content==payload


@pytest.mark.parametrize('name,payload',[('bad.docx',b'not zip'),('bad.xlsx',b'not zip'),('bad.txt',b'\xff\x00'),('old.doc',b'old'),('old.xls',b'old')])
def test_invalid_formats_rejected(admin,kb,name,payload):
    assert admin.post(f'/api/knowledge/{kb}/documents',content=payload,headers={'X-Filename':name}).status_code==400


def test_pdf_preview_and_bounds(admin,kb):
    identifier=indexed(admin,kb)
    result=admin.get(f'/api/documents/{identifier}/preview?page=1')
    assert result.status_code==200,result.text[:200]
    assert result.content.startswith(b'\x89PNG')
    assert admin.get(f'/api/documents/{identifier}/preview?page=0').status_code==400
    assert admin.get(f'/api/documents/{identifier}/preview?page=999').status_code==404


def test_reader_can_search_but_not_write_or_reshare(admin,kb):
    identifier=indexed(admin,kb)
    add_member(admin,kb)
    sign_in(admin,'member','member-password-123')
    assert admin.get('/api/knowledge').json()[0]['permission']=='reader'
    assert search(2,kb,'MikroTik')
    assert admin.get(f'/api/documents/{identifier}/preview').status_code==200
    assert admin.post(f'/api/knowledge/{kb}/documents',content=b'test',headers={'X-Filename':'test.txt'}).status_code==403
    assert admin.post(f'/api/documents/{identifier}/reindex').status_code==403
    assert admin.delete(f'/api/documents/{identifier}').status_code==403
    assert admin.delete(f'/api/knowledge/{kb}').status_code==404
    assert admin.post(f'/api/knowledge/{kb}/members',json={'username':'owner'}).status_code==404
    assert admin.post('/api/assistants',json={'name':'Shared tutor','kb_id':kb}).status_code==201


def test_editor_uploads_count_against_owner_quota(admin,kb,monkeypatch):
    add_member(admin,kb,'editor');sign_in(admin,'member','member-password-123')
    identifier=add_document(admin,kb,'shared.txt',b'Shared MikroTik information')
    assert admin.post(f'/api/documents/{identifier}/reindex').status_code==200
    assert admin.delete(f'/api/documents/{identifier}').status_code==403
    monkeypatch.setattr(settings,'user_storage_mb',0)
    assert admin.post(f'/api/knowledge/{kb}/documents',content=b'test',headers={'X-Filename':'new.txt'}).status_code==413


def test_revocation_blocks_sources_and_retrieval(admin,kb):
    identifier=indexed(admin,kb);add_member(admin,kb)
    admin.delete(f'/api/knowledge/{kb}/members/2')
    sign_in(admin,'member','member-password-123')
    assert admin.get('/api/knowledge').json()==[]
    assert search(2,kb,'MikroTik')==[]
    assert admin.get(f'/api/documents/{identifier}/preview').status_code==404
    assert admin.get(f'/api/documents/{identifier}/download').status_code==404


def test_model_import_checksum_and_no_overwrite(admin):
    payload=gguf();sha=hashlib.sha256(payload).hexdigest()
    bad=admin.post('/api/models/import',content=payload,headers={'X-Filename':'model.gguf','X-SHA256':'0'*64})
    assert bad.status_code==400
    assert not list(settings.models_dir.iterdir())
    result=admin.post('/api/models/import',content=payload,headers={'X-Filename':'model.gguf','X-SHA256':sha})
    assert result.status_code==201,result.text
    assert result.json()['checksum_verified'] and result.json()['sha256']==sha
    assert admin.post('/api/models/import',content=payload,headers={'X-Filename':'model.gguf'}).status_code==409
    assert admin.get('/api/models/inspect?filename=model.gguf').json()['metadata']['general.architecture']=='llama'


def test_model_import_invalid_path_header_and_size(admin,monkeypatch):
    for name in ['../model.gguf','/tmp/model.gguf','..\\model.gguf']:
        assert admin.post('/api/models/import',content=gguf(),headers={'X-Filename':name}).status_code==400
    assert admin.post('/api/models/import',content=b'GGUF',headers={'X-Filename':'bad.gguf'}).status_code==400
    monkeypatch.setattr(settings,'max_model_mb',0)
    assert admin.post('/api/models/import',content=gguf(),headers={'X-Filename':'big.gguf'}).status_code==413
    assert not list(settings.models_dir.iterdir())


def test_profiles_round_trip_and_admin_only(admin,kb):
    (settings.models_dir/'model.gguf').write_bytes(gguf())
    config={'filename':'model.gguf','threads':2,'context':2048}
    assert admin.post('/api/model-profiles',json={'name':'Laptop','config':config}).status_code==201
    saved=admin.get('/api/model-profiles').json()[0]
    assert saved['config']['threads']==2
    add_member(admin,kb);sign_in(admin,'member','member-password-123')
    assert admin.get('/api/model-profiles').status_code==403
    assert admin.post('/api/models/import',content=gguf(),headers={'X-Filename':'x.gguf'}).status_code==403
    assert admin.get('/api/system/check').status_code==403


def test_wizard_checks_and_completion(admin,monkeypatch):
    result=admin.get('/api/system/check')
    assert result.status_code==200 and result.json()['ram_total']>0
    assert admin.post('/api/system/complete').status_code==409
    fake_model(monkeypatch)
    assert admin.post('/api/system/model-test').status_code==200
    assert admin.post('/api/system/complete').status_code==200
    assert db.get_setting('wizard_complete') is True


def test_document_followup_uses_prior_question(admin,kb,monkeypatch):
    indexed(admin,kb);fake_model(monkeypatch)
    captured=[]
    monkeypatch.setattr(runtime,'stream',lambda messages,stop: captured.append(messages) or iter(['Answer [S1]']))
    convo=admin.post('/api/conversations',json={}).json()['id']
    for question in ['What MikroTik training is available?','Explain that further.']:
        result=admin.post(f'/api/conversations/{convo}/chat',json={'content':question,'kb_id':kb})
        assert result.status_code==200
    assert 'MikroTik training' in captured[-1][-1]['content']
    assert 'Explain that further' in captured[-1][-1]['content']


def test_comparison_validates_selection(admin,kb,monkeypatch):
    identifier=indexed(admin,kb);fake_model(monkeypatch)
    convo=admin.post('/api/conversations',json={}).json()['id']
    response=admin.post(f'/api/conversations/{convo}/chat',json={'content':'Compare','kb_id':kb,'task':'compare','document_ids':[identifier]})
    assert response.status_code==400
    response=admin.post(f'/api/conversations/{convo}/chat',json={'content':'Summary','kb_id':kb,'task':'summary','document_ids':[9999]})
    assert response.status_code==404


def test_full_summary_covers_all_chunks(admin,kb,monkeypatch):
    identifier=add_document(admin,kb,'long.txt',('MikroTik training. '+('Different lesson details. '*400)).encode())
    fake_model(monkeypatch)
    calls=[]
    def capture(messages,stop):
        calls.append(messages);return iter(['Concise notes [S1].'])
    monkeypatch.setattr(runtime,'stream',capture)
    convo=admin.post('/api/conversations',json={}).json()['id']
    response=admin.post(f'/api/conversations/{convo}/chat',json={'content':'Summarise all topics','kb_id':kb,'task':'summary','document_ids':[identifier]})
    assert response.status_code==200 and 'All passages processed' in response.text
    stored=admin.get(f'/api/conversations/{convo}/messages').json()[-1]
    assert stored['status']=='complete'
    assert len(stored['sources'])==db.one('SELECT count(*) n FROM chunks WHERE document_id=?',(identifier,))['n']
    for source in stored['sources']:
        assert any('['+source['id']+']' in c[0]['content'] for c in calls)


def test_v1_migration_preserves_data_and_is_repeatable(tmp_path,monkeypatch):
    monkeypatch.setattr(settings,'data_dir',tmp_path)
    monkeypatch.setattr(settings,'models_dir',tmp_path/'models')
    conn=sqlite3.connect(settings.db_path);conn.executescript(db.SCHEMA)
    conn.execute("INSERT INTO users VALUES(1,'old-owner','hash','admin',0)")
    conn.execute("INSERT INTO knowledge_bases VALUES(1,1,'Old PDFs')")
    conn.execute("INSERT INTO documents(kb_id,name,sha256,pdf,size,created) VALUES(1,'old.pdf','checksum',?,4,0)",(b'%PDF',))
    conn.execute('PRAGMA user_version=1');conn.commit();conn.close()
    db.initialize();db.initialize()
    assert db.one('SELECT format,pdf FROM documents')['format']=='pdf'
    assert db.one('SELECT pdf FROM documents')['pdf']==b'%PDF'
    with db.connect() as conn:
        assert conn.execute('PRAGMA user_version').fetchone()[0]==3
        assert not conn.execute('PRAGMA foreign_key_check').fetchall()


def test_comparison_reads_both_documents(admin,kb,monkeypatch):
    first=add_document(admin,kb,'one.txt',b'First service costs 100 naira.')
    second=add_document(admin,kb,'two.txt',b'Second service costs 200 naira.')
    fake_model(monkeypatch)
    convo=admin.post('/api/conversations',json={}).json()['id']
    result=admin.post(f'/api/conversations/{convo}/chat',json={'content':'Compare prices','kb_id':kb,'task':'compare','document_ids':[first,second]})
    assert result.status_code==200
    row=admin.get(f'/api/conversations/{convo}/messages').json()[-1]
    assert row['status']=='complete'
    assert {s['document_id'] for s in row['sources']}=={first,second}


def test_summary_limit_is_explicit(admin,kb,monkeypatch):
    identifier=add_document(admin,kb,'long.txt',b'Long document. '*500)
    fake_model(monkeypatch);monkeypatch.setattr(settings,'max_summary_chunks',1)
    convo=admin.post('/api/conversations',json={}).json()['id']
    result=admin.post(f'/api/conversations/{convo}/chat',json={'content':'Summarise','kb_id':kb,'task':'summary','document_ids':[identifier]})
    assert 'Nothing was silently omitted' in result.text
    assert admin.get(f'/api/conversations/{convo}/messages').json()[-1]['status']=='failed'


def test_office_archive_expansion_limit(admin,kb):
    import zipfile
    payload=io.BytesIO()
    with zipfile.ZipFile(payload,'w',compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr('word/document.xml',b' '*(101*1024**2))
    result=admin.post(f'/api/knowledge/{kb}/documents',content=payload.getvalue(),headers={'X-Filename':'large.docx'})
    assert result.status_code==400 and 'unpacked limit' in result.text
