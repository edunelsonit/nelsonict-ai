"""Bounded extraction; source locations are genuine pages, blocks or row references."""
import csv
import io
import zipfile
from pathlib import Path

FORMATS = {'pdf': 'application/pdf', 'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
           'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
           'txt': 'text/plain', 'md': 'text/markdown', 'csv': 'text/csv'}


def validate(payload, name):
    kind = Path(name).suffix.lower().lstrip('.')
    if kind not in FORMATS:
        raise ValueError('Supported formats: PDF, DOCX, TXT, Markdown, CSV and XLSX. Convert legacy DOC/XLS first.')
    if kind == 'pdf' and not payload.startswith(b'%PDF-'):
        raise ValueError('File is not a PDF.')
    if kind in ('docx', 'xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(x.file_size for x in entries) > 100 * 1024**2:
                    raise ValueError('Office document exceeds the 100 MB unpacked limit.')
                required = 'word/document.xml' if kind == 'docx' else 'xl/workbook.xml'
                if required not in archive.namelist():
                    raise ValueError('File contents do not match its Office extension.')
        except zipfile.BadZipFile:
            raise ValueError('Invalid Office document.')
    if kind in ('txt', 'md', 'csv'):
        try:
            text = payload.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise ValueError('Save text/CSV documents in UTF-8 encoding.')
        if '\x00' in text:
            raise ValueError('Binary contents are not supported as text.')
    return kind


def extract_units(payload, kind):
    """Yield (unit_number, text, location). PDF OCR remains in index_document."""
    if kind in ('txt', 'md'):
        text = payload.decode('utf-8-sig')
        for i in range(0, len(text), 1200):
            start = text.count('\n', 0, i) + 1
            end = start + text[i:i+1200].count('\n')
            yield i // 1200 + 1, text[i:i+1200], f'Lines {start}–{end}'
    elif kind == 'csv':
        csv.field_size_limit(1024 * 1024)
        reader = csv.reader(io.StringIO(payload.decode('utf-8-sig')))
        header = next(reader, [])
        if len(header) > 200:
            raise ValueError('CSV exceeds 200 columns.')
        for rownum, row in enumerate(reader, 2):
            if rownum > 50001 or len(row) > 200:
                raise ValueError('Table exceeds 50,000 rows or 200 columns.')
            text = '; '.join(f'{header[i] if i < len(header) else "Column " + str(i+1)}: {v}' for i,v in enumerate(row))
            yield rownum-1, text, f'CSV row {rownum}'
    elif kind == 'docx':
        from docx import Document
        from docx.table import Table
        doc = Document(io.BytesIO(payload))
        for number, block in enumerate(doc.iter_inner_content(), 1):
            if isinstance(block, Table):
                text = '\n'.join(' | '.join(cell.text for cell in row.cells) for row in block.rows)
                label = f'Table/block {number}'
            else:
                text, label = block.text, f'Paragraph/block {number}'
            yield number, text, label
    elif kind == 'xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
        number = 0
        try:
            if len(book.worksheets) > 100:
                raise ValueError('Workbook exceeds 100 sheets.')
            for sheet in book.worksheets:
                if (sheet.max_column or 0) > 200 or (sheet.max_row or 0) > 50000:
                    raise ValueError('Sheet exceeds 50,000 rows or 200 columns.')
                for rownum, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    number += 1
                    if number > 50000:
                        raise ValueError('Workbook exceeds 50,000 total rows.')
                    values = ['' if x is None else str(x) for x in row]
                    if any(values):
                        yield number, f'Sheet {sheet.title}; row {rownum}: ' + ' | '.join(values), f'{sheet.title}!row {rownum}'
        finally:
            book.close()
